#!/usr/bin/env python3
"""Client subscription tracker + branded PDF invoice.

Generalised from the PedalAI/Clarity i2 build so it works for any client.
Everything client-specific lives in the workbook, not in this file.

  init      create a tracker workbook for a client
  invoice   render a PDF for one month, reading the workbook

The workbook is the single source of truth. `invoice` only READS it, so
regenerating can never clobber figures you typed. A blank amount is never
billed and is always reported — you cannot silently under-invoice by
forgetting a line.

Requires: openpyxl, and Chrome (headless) for PDF rendering.

  python3 invoice.py init    --client "Acme Co" --out ~/Acme
  python3 invoice.py invoice --book ~/Acme/Acme_Subscriptions.xlsx [--month Sep-2026]
"""
from __future__ import annotations

import argparse
import base64
import html
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl required:  pip install openpyxl")

CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
]

HEADERS = ["Vendor", "Category", "Plan / Tier", "Billing basis", "Status",
           "Verified", "Invoice description"]
FIRST_MONTH_COL = len(HEADERS) + 1
HEAD_ROW = 4
MNAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

INK, MONEY = "1F2A30", '"$"#,##0.00'
HEAD_FILL = PatternFill("solid", fgColor="404040")   # matches the invoice table header
BAND = PatternFill("solid", fgColor="F2F4F5")
NEEDS = PatternFill("solid", fgColor="FFF4D6")       # amber = you must enter this
THIN = Side(style="thin", color="D5DADD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def month_label(y: int, m: int) -> str:
    return f"{MNAMES[m - 1]}-{y}"


def find_chrome() -> str:
    for c in CHROME_CANDIDATES:
        if Path(c).exists():
            return c
    sys.exit("Chrome/Chromium not found — needed to render the PDF.")


# ── init ────────────────────────────────────────────────────────────────

def cmd_init(a) -> None:
    out_dir = Path(a.out).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    slug = "".join(ch for ch in a.client if ch.isalnum() or ch in " -_").strip().replace(" ", "_")
    book = out_dir / f"{slug}_Subscriptions.xlsx"
    if book.exists() and not a.force:
        sys.exit(f"{book} already exists — refusing to overwrite. Use --force to replace it.")

    start_y, start_m = (int(x) for x in a.start.split("-")) if a.start else (date.today().year, date.today().month)
    months = []
    y, m = start_y, start_m
    for _ in range(a.months):
        months.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Subscriptions"
    ws["A1"] = f"{a.client} — Platform Subscriptions"
    ws["A1"].font = Font(bold=True, size=15, color=INK)
    ws["A2"] = ("Enter the amount actually billed in each month's column. Amber cells still need a real figure. "
                "The TOTAL row and the PDF invoice both read from this sheet.")
    ws["A2"].font = Font(size=9, italic=True, color="6B767B")

    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=HEAD_ROW, column=i, value=h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=10), HEAD_FILL, BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for j, (yy, mm) in enumerate(months):
        c = ws.cell(row=HEAD_ROW, column=FIRST_MONTH_COL + j, value=month_label(yy, mm))
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=10), HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center")

    # One example row so the shape is obvious, then blanks to fill in.
    seed = [("(example) Apollo.io", "Data", "Organization seat", "Monthly subscription",
             "Active", "Needs check", "Apollo.io — B2B contact and company database")]
    rows = seed + [("", "", "", "", "", "ENTER AMOUNT", "")] * max(0, a.rows - 1)
    for r, v in enumerate(rows, start=HEAD_ROW + 1):
        for i, val in enumerate(v, start=1):
            c = ws.cell(row=r, column=i, value=val or None)
            c.border, c.font = BOX, Font(size=10, color=INK)
            c.alignment = Alignment(vertical="center", wrap_text=(i == 7))
            if r % 2 == 0:
                c.fill = BAND
        for j in range(len(months)):
            c = ws.cell(row=r, column=FIRST_MONTH_COL + j)
            c.number_format, c.border = MONEY, BOX
            c.alignment = Alignment(horizontal="right")
            c.fill = NEEDS

    trow = HEAD_ROW + len(rows) + 1
    ws.cell(row=trow, column=1, value="TOTAL (USD)").font = Font(bold=True, size=11, color=INK)
    ws.cell(row=trow, column=1).border = BOX
    for i in range(2, FIRST_MONTH_COL):
        ws.cell(row=trow, column=i).border = BOX
    for j in range(len(months)):
        col = get_column_letter(FIRST_MONTH_COL + j)
        c = ws.cell(row=trow, column=FIRST_MONTH_COL + j,
                    value=f"=SUM({col}{HEAD_ROW+1}:{col}{trow-1})")
        c.number_format, c.border = MONEY, BOX
        c.font = Font(bold=True, size=11, color=INK)
        c.fill = PatternFill("solid", fgColor="E3EAEF")
        c.alignment = Alignment(horizontal="right")

    for col, w in {"A": 21, "B": 14, "C": 21, "D": 21, "E": 17, "F": 19, "G": 62}.items():
        ws.column_dimensions[col].width = w
    for j in range(len(months)):
        ws.column_dimensions[get_column_letter(FIRST_MONTH_COL + j)].width = 12
    ws.row_dimensions[HEAD_ROW].height = 26
    ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=FIRST_MONTH_COL)

    iv = wb.create_sheet("Invoice")
    iv["A1"] = "Invoice settings"
    iv["A1"].font = Font(bold=True, size=15, color=INK)
    iv["A2"] = ("Read by the PDF generator. Change 'Billing month' to invoice a different period. "
                "'From' fields are YOUR company; 'Bill to' is the client.")
    iv["A2"].font = Font(size=9, italic=True, color="6B767B")
    fields = [
        ("Invoice number", 1),
        ("Invoice date", date.today().strftime("%B %d, %Y").replace(" 0", " ")),
        ("Payment due", ""),
        ("Billing month", month_label(*months[0])),
        ("Terms (days)", 10),
        ("", ""),
        ("From — company", a.company or ""),
        ("From — country", "United States"),
        ("From — logo path", a.logo or ""),
        ("From — footer", f"{a.company} © {date.today().year}" if a.company else ""),
        ("", ""),
        ("Bill to — company", a.client),
        ("Bill to — contact", ""),
        ("Bill to — street", ""),
        ("Bill to — suite", ""),
        ("Bill to — city/state/zip", ""),
        ("Bill to — country", "United States"),
        ("Bill to — phone", ""),
        ("Bill to — email", ""),
        ("", ""),
        ("Include zero-value lines", "yes"),
    ]
    for r, (k, v) in enumerate(fields, start=4):
        if not k:
            continue
        ka = iv.cell(row=r, column=1, value=k)
        ka.font, ka.border = Font(bold=True, size=10, color=INK), BOX
        vb = iv.cell(row=r, column=2, value=v)
        vb.font, vb.border = Font(size=10, color=INK), BOX
        if v == "":
            vb.fill = NEEDS
    iv.column_dimensions["A"].width = 26
    iv.column_dimensions["B"].width = 46

    wb.save(book)
    print(f"wrote {book}")
    print(f"  {len(rows)} vendor rows · {len(months)} months "
          f"({month_label(*months[0])} .. {month_label(*months[-1])})")
    print("  next: fill the vendor rows + the amber Invoice cells, then run `invoice`")


# ── invoice ─────────────────────────────────────────────────────────────

def read_book(book: Path, month_override: str | None):
    wb = load_workbook(book, data_only=True)
    iv = wb["Invoice"]
    cfg = {}
    for r in range(4, iv.max_row + 1):
        k = iv.cell(row=r, column=1).value
        if k:
            cfg[str(k).strip()] = iv.cell(row=r, column=2).value

    month = month_override or str(cfg.get("Billing month", "")).strip()
    ws = wb["Subscriptions"]
    col = None
    for c in range(FIRST_MONTH_COL, ws.max_column + 1):
        if str(ws.cell(row=HEAD_ROW, column=c).value).strip() == month:
            col = c
            break
    if col is None:
        avail = [str(ws.cell(row=HEAD_ROW, column=c).value)
                 for c in range(FIRST_MONTH_COL, ws.max_column + 1)]
        sys.exit(f"Month {month!r} not found. Available: {', '.join(avail)}")

    include_zero = str(cfg.get("Include zero-value lines", "yes")).strip().lower() in ("yes", "y", "true")
    items, total, missing = [], 0.0, []
    for r in range(HEAD_ROW + 1, ws.max_row + 1):
        vendor = ws.cell(row=r, column=1).value
        if not vendor or str(vendor).strip().upper().startswith("TOTAL"):
            continue
        status = str(ws.cell(row=r, column=5).value or "")
        desc = ws.cell(row=r, column=7).value or vendor
        amt = ws.cell(row=r, column=col).value
        if amt is None:
            missing.append(str(vendor))
            continue                       # never invent a figure
        amt = float(amt)
        if status.lower().startswith("cancelled") and amt == 0:
            continue                       # retired vendor, nothing to bill
        if amt == 0 and not include_zero:
            continue
        items.append((str(desc), amt))
        total += amt
    return cfg, month, items, total, missing


def render_html(cfg, month, items, total) -> str:
    def g(k, d=""):
        v = cfg.get(k)
        return html.escape(str(v)) if v not in (None, "") else d

    logo_path = str(cfg.get("From — logo path") or "").strip()
    logo_tag = ""
    if logo_path and Path(logo_path).expanduser().exists():
        b64 = base64.b64encode(Path(logo_path).expanduser().read_bytes()).decode()
        logo_tag = f'<img src="data:image/png;base64,{b64}" alt="">'
    else:
        logo_tag = f'<div class="wordmark">{g("From — company")}</div>'

    addr = [g("Bill to — company"), g("Bill to — contact"), g("Bill to — street"),
            g("Bill to — suite"), g("Bill to — city/state/zip"), g("Bill to — country"),
            g("Bill to — phone"), g("Bill to — email")]
    rows = "".join(
        f"<tr><td class='d'>{html.escape(d)}</td><td class='q'>1</td>"
        f"<td class='m'>${a:,.2f}</td><td class='m'>${a:,.2f}</td></tr>" for d, a in items)

    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
  @page {{ size: Letter; margin: 0.6in 0.7in; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: Arial, Helvetica, sans-serif; color:#1a1a1a; font-size:10.5pt; margin:0; }}
  .top {{ display:flex; align-items:flex-start; justify-content:space-between; gap:24px; margin-bottom:26px; }}
  .top img {{ width:3.1in; }}
  .wordmark {{ font-size:20pt; font-weight:bold; }}
  .title {{ text-align:right; }}
  .title .word {{ font-size:28pt; font-weight:bold; line-height:1; }}
  .mid {{ display:flex; justify-content:space-between; gap:30px; margin-bottom:26px; }}
  .billto {{ line-height:1.45; }} .billto .lbl {{ font-weight:bold; }}
  .meta {{ text-align:right; line-height:1.6; white-space:nowrap; }} .meta .lbl {{ font-weight:bold; }}
  table {{ width:100%; border-collapse:collapse; }}
  thead th {{ background:#404040; color:#fff; font-weight:bold; font-size:10pt; padding:7px 9px; text-align:left; }}
  thead th.q {{ text-align:center; }} thead th.m {{ text-align:right; }}
  tbody td {{ padding:7px 9px; border-bottom:1px solid #d9d9d9; font-size:10pt; vertical-align:top; }}
  tbody tr:nth-child(odd) td {{ background:#f2f2f2; }}
  td.q {{ text-align:center; white-space:nowrap; }} td.m {{ text-align:right; white-space:nowrap; }}
  tfoot td {{ padding:9px; font-weight:bold; font-size:11pt; border-top:2px solid #404040; }}
  tfoot td.m {{ text-align:right; }}
  .foot {{ position:fixed; bottom:0; left:0; right:0; text-align:center; font-size:8.5pt; color:#666; }}
</style></head><body>
  <div class="top">{logo_tag}
    <div class="title"><div class="word">INVOICE</div><div>{g("From — country")}</div></div>
  </div>
  <div class="mid">
    <div class="billto"><span class="lbl">Bill to:</span><br>{"<br>".join(a for a in addr if a)}</div>
    <div class="meta">
      <span class="lbl">Invoice Number:</span> {g("Invoice number")}<br>
      <span class="lbl">Invoice Date:</span> {g("Invoice date")}<br>
      <span class="lbl">Payment Due:</span> {g("Payment due")}<br>
      <span class="lbl">Billing Period:</span> {html.escape(month)}<br>
      <span class="lbl">Amount Due (USD):</span> ${total:,.2f}
    </div>
  </div>
  <table>
    <thead><tr><th>Items</th><th class="q">Quantity</th><th class="m">Price</th><th class="m">Amount</th></tr></thead>
    <tbody>{rows}</tbody>
    <tfoot><tr><td colspan="3">TOTAL DUE (USD)</td><td class="m">${total:,.2f}</td></tr></tfoot>
  </table>
  <div class="foot">{g("From — footer")}</div>
</body></html>"""


def cmd_invoice(a) -> None:
    book = Path(a.book).expanduser()
    if not book.exists():
        sys.exit(f"{book} not found — run `init` first.")
    cfg, month, items, total, missing = read_book(book, a.month)
    if not items:
        sys.exit(f"No billable line items for {month}.")

    slug = str(cfg.get("From — company") or "Invoice").replace(" ", "_")
    out = Path(a.out).expanduser() if a.out else book.parent / f"{slug}_Invoice_{month}.pdf"

    tmp = Path(tempfile.mkdtemp())
    (tmp / "invoice.html").write_text(render_html(cfg, month, items, total), encoding="utf-8")
    subprocess.run([find_chrome(), "--headless", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={out}", str(tmp / "invoice.html")],
                   check=True, capture_output=True, timeout=180)
    shutil.rmtree(tmp, ignore_errors=True)

    print(f"wrote {out}")
    print(f"period {month} · {len(items)} line items · total ${total:,.2f}")
    if missing:
        # Loud on purpose: a blank cell is the one way to under-bill silently.
        print(f"NOT BILLED (blank in the sheet): {', '.join(missing)}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    i = sub.add_parser("init", help="create a tracker workbook for a client")
    i.add_argument("--client", required=True, help="who you're billing")
    i.add_argument("--company", default="", help="your company (appears on the invoice)")
    i.add_argument("--out", required=True, help="directory for the workbook")
    i.add_argument("--logo", default="", help="path to your logo PNG")
    i.add_argument("--months", type=int, default=12)
    i.add_argument("--rows", type=int, default=12)
    i.add_argument("--start", default="", help="YYYY-MM (default: this month)")
    i.add_argument("--force", action="store_true", help="overwrite an existing workbook")
    i.set_defaults(func=cmd_init)

    v = sub.add_parser("invoice", help="render the PDF for one month")
    v.add_argument("--book", required=True)
    v.add_argument("--month", help="e.g. Sep-2026 (default: the workbook's Billing month)")
    v.add_argument("--out")
    v.set_defaults(func=cmd_invoice)

    a = ap.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
